import csv
from pathlib import Path
import openpyxl
import pytest

from app.database.models import Business, Finding, LeadScore, LighthouseResult, PipelineRun, WebsiteProfile
from app.exports.csv_export import export_leads_to_csv
from app.exports.excel_export import export_leads_to_excel


@pytest.fixture
def sample_lead_data():
    b = Business(
        id=10,
        name="Lone Star Steakhouse",
        normalized_name="lone star steakhouse",
        vertical="restaurant",
        city="Austin",
        state="Texas",
        country="United States",
        website="http://lonestarsteak.example",
        phone="512-555-4321",
        public_email="hello@lonestarsteak.example",
    )
    score = LeadScore(
        business_id=10,
        total_score=85.0,
        tier="A+",
        website_score=45.0,
        seo_opportunity_score=18.0,
        vertical_opportunity_score=14.0,
        contactability_score=8.0,
        recommended_service="WEBSITE REDESIGN + LOCAL SEO",
        recommended_angle="Modernize online reservation funnel.",
    )
    profile = WebsiteProfile(
        business_id=10,
        requested_url="http://lonestarsteak.example",
        reachable=True,
        has_viewport=False,
    )
    lh = LighthouseResult(
        business_id=10,
        performance_score=35.0,
        seo_score=52.0,
    )
    findings = [
        Finding(
            business_id=10,
            finding_code="NO_VIEWPORT",
            category="WEBSITE",
            severity="HIGH",
            title="Missing Mobile Viewport",
            description="Site is not mobile optimized.",
            score_points=20,
        )
    ]
    return [(b, score, profile, lh, findings)]


def test_csv_export(tmp_path, sample_lead_data):
    csv_file = tmp_path / "test_leads.csv"
    export_leads_to_csv(sample_lead_data, csv_file)

    assert csv_file.exists()
    with open(csv_file, "r", encoding="utf-8") as f:
        reader = list(csv.reader(f))
        assert len(reader) == 2  # header + 1 row
        header = reader[0]
        row = reader[1]
        assert "Lone Star Steakhouse" in row
        assert "A+" in row
        assert "WEBSITE REDESIGN + LOCAL SEO" in row


def test_excel_export(tmp_path, sample_lead_data):
    xlsx_file = tmp_path / "test_leads.xlsx"
    run_info = PipelineRun(
        id=1,
        market="Austin, Texas, USA",
        vertical="restaurant",
        target=10,
        raw_found=20,
        deduped=18,
        filtered=15,
        websites_processed=12,
        qualified=1,
        exported=1,
    )

    b, score, profile, lh, findings = sample_lead_data[0]
    export_leads_to_excel(
        run_info=run_info,
        qualified_leads=sample_lead_data,
        all_discovered=[b],
        rejected_businesses=[],
        output_path=xlsx_file,
    )

    assert xlsx_file.exists()

    # Open workbook and verify tabs
    wb = openpyxl.load_workbook(xlsx_file)
    expected_sheets = [
        "HOT LEADS",
        "A LEADS",
        "B LEADS",
        "NO WEBSITE",
        "WEBSITE REDESIGN",
        "SEO OPPORTUNITIES",
        "ALL QUALIFIED",
        "ALL DISCOVERED",
        "REJECTED",
        "RUN SUMMARY",
    ]
    for s in expected_sheets:
        assert s in wb.sheetnames

    # Check HOT LEADS has our lead
    ws_hot = wb["HOT LEADS"]
    assert ws_hot.max_row == 2
    assert ws_hot.cell(row=2, column=4).value == "Lone Star Steakhouse"
    assert ws_hot.cell(row=2, column=3).value == "A+"

    # Check RUN SUMMARY
    ws_sum = wb["RUN SUMMARY"]
    assert ws_sum.cell(row=3, column=2).value == "Austin, Texas, USA"
