import logging
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.database.models import Business, Finding, LeadScore, LighthouseResult, PipelineRun, WebsiteProfile

logger = logging.getLogger("leadengine.exports.excel")

HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
HEADER_FONT = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
CELL_FONT = Font(name="Segoe UI", size=9)
BOLD_FONT = Font(name="Segoe UI", size=9, bold=True)

THIN_BORDER = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)

TIER_COLORS = {
    "A+": ("DCFCE7", "166534"),  # Light green bg, dark green text
    "A": ("E0F2FE", "075985"),   # Light blue bg, dark blue text
    "B": ("FEF3C7", "92400E"),   # Light yellow bg, dark amber text
    "C": ("FEE2E2", "991B1B"),   # Light red bg, dark red text
    "D": ("F1F5F9", "475569"),   # Slate gray
}

COLUMNS = [
    ("Rank", 7, "center"),
    ("Score", 8, "center"),
    ("Tier", 7, "center"),
    ("Business", 28, "left"),
    ("Vertical", 14, "left"),
    ("City", 16, "left"),
    ("State", 12, "left"),
    ("Country", 14, "left"),
    ("Website", 28, "left"),
    ("Website Status", 15, "center"),
    ("Email", 26, "left"),
    ("Email Status", 13, "center"),
    ("Phone", 18, "left"),
    ("Contact Form", 24, "left"),
    ("Instagram", 22, "left"),
    ("Facebook", 22, "left"),
    ("LinkedIn", 22, "left"),
    ("Performance", 13, "center"),
    ("SEO", 10, "center"),
    ("Accessibility", 13, "center"),
    ("Best Practices", 13, "center"),
    ("Main Problems", 42, "left"),
    ("Recommended Service", 32, "left"),
    ("Sales Angle", 48, "left"),
    ("Technologies", 24, "left"),
    ("Schema Types", 22, "left"),
    ("Source", 14, "center"),
    ("Confidence", 12, "center"),
    ("Audit Date", 16, "center"),
]


def _clean_val(val: Any) -> str:
    if val is None:
        return ""
    s = str(val).strip()
    if s.lower() in ("nan", "none", "null"):
        return ""
    return s


def _format_lead_row(
    rank: int,
    business: Business,
    score: Optional[LeadScore],
    profile: Optional[WebsiteProfile],
    lh: Optional[LighthouseResult],
    findings: List[Finding],
) -> List[Any]:
    if not business.website:
        web_status = "NO_WEBSITE"
    elif profile and profile.reachable:
        web_status = "REACHABLE"
    else:
        web_status = "UNREACHABLE"

    email = _clean_val(business.public_email)
    email_status = "FOUND" if email else "NONE"
    phone = _clean_val(business.normalized_phone or business.phone)
    top_problems = " | ".join(f.title for f in findings[:5]) if findings else "None detected"
    techs = " | ".join(profile.detected_technologies) if profile and profile.detected_technologies else ""
    schemas = " | ".join(profile.schema_types) if profile and profile.schema_types else ""

    sources = []
    if business.source_overture:
        sources.append("Overture")
    if business.source_osm:
        sources.append("OSM")
    source_str = " + ".join(sources) if sources else "Direct"

    return [
        rank,
        score.total_score if score else "",
        score.tier if score else "",
        business.name,
        business.vertical,
        _clean_val(business.city),
        _clean_val(business.state),
        _clean_val(business.country),
        _clean_val(business.website),
        web_status,
        email,
        email_status,
        phone,
        profile.contact_page_url if profile and profile.contact_page_url else "",
        _clean_val(business.instagram),
        _clean_val(business.facebook),
        _clean_val(business.linkedin),
        lh.performance_score if lh and lh.performance_score is not None else "",
        lh.seo_score if lh and lh.seo_score is not None else "",
        lh.accessibility_score if lh and lh.accessibility_score is not None else "",
        lh.best_practices_score if lh and lh.best_practices_score is not None else "",
        top_problems,
        score.recommended_service if score else "",
        score.recommended_angle if score else "",
        techs,
        schemas,
        source_str,
        f"{business.source_confidence:.2f}",
        score.calculated_at if score and score.calculated_at else (business.updated_at or ""),
    ]


def _populate_sheet(ws: openpyxl.worksheet.worksheet.Worksheet, rows: List[List[Any]]) -> None:
    # Write Headers
    for col_idx, (col_name, col_width, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = col_width

    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"

    for r_idx, row_data in enumerate(rows, start=2):
        ws.row_dimensions[r_idx].height = 20
        tier_val = str(row_data[2])
        for c_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            align = COLUMNS[c_idx - 1][2]
            cell.alignment = Alignment(horizontal=align, vertical="center")
            cell.font = CELL_FONT
            cell.border = THIN_BORDER

            # Tier styling
            if c_idx == 3 and tier_val in TIER_COLORS:
                bg, fg = TIER_COLORS[tier_val]
                cell.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
                cell.font = Font(name="Segoe UI", size=9, bold=True, color=fg)

    if len(rows) > 0:
        last_col = get_column_letter(len(COLUMNS))
        ws.auto_filter.ref = f"A1:{last_col}{len(rows) + 1}"


def export_leads_to_excel(
    run_info: PipelineRun,
    qualified_leads: List[Tuple[Business, LeadScore, Optional[WebsiteProfile], Optional[LighthouseResult], List[Finding]]],
    all_discovered: List[Business],
    rejected_businesses: List[Business],
    output_path: Path,
) -> Path:
    """Generate an Excel workbook with segregated tabs, conditional badges, and run summary."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # 1. Hot Leads (Tier A+)
    ws_hot = wb.create_sheet(title="HOT LEADS")
    hot_rows = [
        _format_lead_row(i, b, s, p, lh, f)
        for i, (b, s, p, lh, f) in enumerate(qualified_leads, start=1)
        if s.tier == "A+"
    ]
    _populate_sheet(ws_hot, hot_rows)

    # 2. A Leads (Tier A)
    ws_a = wb.create_sheet(title="A LEADS")
    a_rows = [
        _format_lead_row(i, b, s, p, lh, f)
        for i, (b, s, p, lh, f) in enumerate(qualified_leads, start=1)
        if s.tier == "A"
    ]
    _populate_sheet(ws_a, a_rows)

    # 3. B Leads (Tier B)
    ws_b = wb.create_sheet(title="B LEADS")
    b_rows = [
        _format_lead_row(i, b, s, p, lh, f)
        for i, (b, s, p, lh, f) in enumerate(qualified_leads, start=1)
        if s.tier == "B"
    ]
    _populate_sheet(ws_b, b_rows)

    # 4. No Website Leads
    ws_noweb = wb.create_sheet(title="NO WEBSITE")
    noweb_rows = [
        _format_lead_row(i, b, s, p, lh, f)
        for i, (b, s, p, lh, f) in enumerate(qualified_leads, start=1)
        if not b.website
    ]
    _populate_sheet(ws_noweb, noweb_rows)

    # 5. Website Redesign
    ws_redesign = wb.create_sheet(title="WEBSITE REDESIGN")
    redesign_rows = [
        _format_lead_row(i, b, s, p, lh, f)
        for i, (b, s, p, lh, f) in enumerate(qualified_leads, start=1)
        if "REDESIGN" in s.recommended_service or "REBUILD" in s.recommended_service or "RESERVATION" in s.recommended_service
    ]
    _populate_sheet(ws_redesign, redesign_rows)

    # 6. SEO Opportunities
    ws_seo = wb.create_sheet(title="SEO OPPORTUNITIES")
    seo_rows = [
        _format_lead_row(i, b, s, p, lh, f)
        for i, (b, s, p, lh, f) in enumerate(qualified_leads, start=1)
        if "SEO" in s.recommended_service or (lh and lh.seo_score and lh.seo_score < 65)
    ]
    _populate_sheet(ws_seo, seo_rows)

    # 7. All Qualified Leads
    ws_all_qual = wb.create_sheet(title="ALL QUALIFIED")
    all_qual_rows = [
        _format_lead_row(i, b, s, p, lh, f)
        for i, (b, s, p, lh, f) in enumerate(qualified_leads, start=1)
    ]
    _populate_sheet(ws_all_qual, all_qual_rows)

    # 8. All Discovered
    ws_all_disc = wb.create_sheet(title="ALL DISCOVERED")
    disc_rows = [
        _format_lead_row(i, b, None, None, None, [])
        for i, b in enumerate(all_discovered, start=1)
    ]
    _populate_sheet(ws_all_disc, disc_rows)

    # 9. Rejected / Chains
    ws_rej = wb.create_sheet(title="REJECTED")
    rej_rows = [
        _format_lead_row(i, b, None, None, None, [])
        for i, b in enumerate(rejected_businesses, start=1)
    ]
    _populate_sheet(ws_rej, rej_rows)

    # 10. Run Summary Sheet
    ws_summary = wb.create_sheet(title="RUN SUMMARY")
    ws_summary.column_dimensions["A"].width = 28
    ws_summary.column_dimensions["B"].width = 36

    # Summary Title
    t_cell = ws_summary.cell(row=1, column=1, value="LEADENGINE PROSPECTING SUMMARY")
    t_cell.font = Font(name="Segoe UI", size=14, bold=True, color="1E293B")

    stats = [
        ("Market", run_info.market),
        ("Vertical", run_info.vertical),
        ("Target Requested", run_info.target),
        ("Raw Discovered", run_info.raw_found),
        ("After Deduplication", run_info.deduped),
        ("Filtered Candidates", run_info.filtered),
        ("Websites Audited", run_info.websites_processed),
        ("Total Qualified Leads", run_info.qualified),
        ("Exported Leads", run_info.exported),
        ("A+ Leads (Hot)", len(hot_rows)),
        ("A Leads", len(a_rows)),
        ("B Leads", len(b_rows)),
        ("No-Website Opportunities", len(noweb_rows)),
        ("Run Status", run_info.status),
        ("Started At", run_info.started_at or ""),
        ("Finished At", run_info.finished_at or ""),
        ("Output File", str(output_path.name)),
    ]

    for idx, (label, val) in enumerate(stats, start=3):
        lbl_cell = ws_summary.cell(row=idx, column=1, value=label)
        lbl_cell.font = BOLD_FONT
        lbl_cell.border = THIN_BORDER
        val_cell = ws_summary.cell(row=idx, column=2, value=val)
        val_cell.font = CELL_FONT
        val_cell.border = THIN_BORDER

    wb.save(output_path)
    logger.info(f"Generated comprehensive Excel workbook: {output_path}")
    return output_path
