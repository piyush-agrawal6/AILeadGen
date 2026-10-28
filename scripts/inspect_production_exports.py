"""Read the actual downloaded acceptance files, without modifying them."""
import csv
import json
from pathlib import Path
from openpyxl import load_workbook
from bs4 import BeautifulSoup

root=Path('output/production-acceptance')
with (root/'leads.csv').open(encoding='utf-8-sig',newline='') as file:
    reader=csv.DictReader(file);headers=reader.fieldnames;rows=list(reader)
required=['Website Opportunity','SEO Opportunity','Security Opportunity','Primary Opportunity','Secondary Opportunity','Business Scale','Contactability','Lead Status','Assigned User','First Seen','Last Seen']
assert len(rows)==5 and all(field in headers for field in required)
for row in rows:
    for name in required[:3]:
        if row[name]:assert 0<=float(row[name])<=100
workbook=load_workbook(root/'leads.xlsx',read_only=True,data_only=True)
sheet=list(workbook['ALL QUALIFIED'].values)
assert any(all(field in row for field in required) for row in sheet)
soup=BeautifulSoup((root/'report.html').read_text(encoding='utf-8'),'html.parser')
text=soup.get_text(' ',strip=True)
assert 'Service opportunities' in text and all(name in text for name in ['Website','SEO','Security'])
assert not soup.find('script')
result={'csv':{'rows':len(rows),'opportunity_and_workflow_columns':required,'business_names':[row['Business'] for row in rows]},'xlsx':{'sheets':workbook.sheetnames,'phase3_headers_present':True},'html':{'service_opportunities_present':True,'no_scripts':True,'bytes':(root/'report.html').stat().st_size}}
workbook.close()
Path('docs/production-export-inspection.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
print(json.dumps(result,indent=2))
