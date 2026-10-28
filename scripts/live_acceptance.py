"""Explicit small production acceptance against the signed internal API."""
import argparse
import hashlib
import hmac
import json
import sys
import time
import uuid
import csv
import io
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import httpx
from app.config.settings import get_settings
from app.api.security import canonical_request

state_path=Path('docs/live-acceptance.json')
def call(method,path,data=None):
    settings=get_settings();body=json.dumps(data,separators=(',',':')).encode() if data is not None else b''
    stamp=str(int(time.time()));nonce=uuid.uuid4().hex;actor=settings.allowed_emails[0]
    signature=hmac.new(settings.api_signing_secret.encode(),canonical_request(stamp,nonce,method,path,body,actor).encode(),hashlib.sha256).hexdigest()
    headers={'x-le-timestamp':stamp,'x-le-nonce':nonce,'x-le-user':actor,'x-le-signature':signature,'Content-Type':'application/json'}
    url=Path('deploy/api-url.txt').read_text(encoding='utf-8-sig').strip()
    response=httpx.request(method,url+path,content=body,headers=headers,timeout=120)
    if response.status_code>=400:raise RuntimeError(f'{method} {path}: HTTP {response.status_code}: {response.text[:400]}')
    return response

def main(action):
    if action=='start':
        assert not state_path.exists(),'Existing acceptance record: inspect it instead of starting another scan.'
        payload={'cities':['Austin, Texas, United States'],'verticals':['dentists'],'target':5,'opportunities':['WEBSITE','SEO','SECURITY'],'quality':'balanced','results':'include_previous','lighthouse_enabled':True,'request_id':str(uuid.uuid4())}
        campaign=call('POST','/api/campaigns',payload).json()
        state={'campaign_id':campaign['id'],'configuration':payload,'created_at':campaign['created_at'],'internal_signed_api':True,'google_oauth_verified':False}
    else:state=json.loads(state_path.read_text())
    if action in ('cancel','resume'):
        call('POST','/api/campaigns/'+state['campaign_id']+'/'+action)
        state.setdefault('lifecycle_actions',[]).append({'action':action,'at':time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime())})
    campaign=call('GET','/api/campaigns/'+state['campaign_id']).json()
    state['latest_campaign']=campaign
    state['checked_at']=time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime())
    if campaign['status']=='COMPLETED':
        leads=call('GET','/api/intelligence/leads?campaign='+state['campaign_id']).json()
        state['results']=leads
        state['events_count']=len(call('GET',f"/api/scans/{campaign['runs'][0]['id']}/events?limit=1000").json())
        if action=='verify':
            assert len(leads['items'])>=5,'Completed scan returned fewer than five qualified leads'
            lead=leads['items'][0];lid=lead['id'];actor=get_settings().allowed_emails[0]
            changed=call('PATCH',f'/api/leads/{lid}/workflow',{'lead_status':'SHORTLISTED','assigned_to':actor}).json()
            assert changed['lead_status']=='SHORTLISTED' and changed['assigned_to']==actor
            note='Production acceptance: report and workflow reviewed; no outreach sent.'
            notes=call('POST',f'/api/leads/{lid}/notes',{'body':note}).json()
            assert any(item['body']==note for item in notes['notes'])
            changed=call('PATCH',f'/api/leads/{lid}/workflow',{'lead_status':'CONTACTED'}).json()
            assert changed['lead_status']=='CONTACTED'
            refreshed=call('GET',f'/api/leads/{lid}/workspace').json()
            assert refreshed['lead_status']=='CONTACTED' and refreshed['assigned_to']==actor
            output=Path('output/production-acceptance');output.mkdir(parents=True,exist_ok=True)
            downloads={}
            for kind in ('xlsx','csv'):
                response=call('GET',f"/api/scans/{campaign['runs'][0]['id']}/export/{kind}")
                assert len(response.content)>50
                (output/('leads.'+kind)).write_bytes(response.content)
                downloads[kind]={'status':response.status_code,'bytes':len(response.content),'sha256':hashlib.sha256(response.content).hexdigest()}
                if kind=='csv':
                    rows=list(csv.DictReader(io.StringIO(response.text)))
                    assert len(rows)==5 and all(name in rows[0] for name in ['Website Opportunity','SEO Opportunity','Security Opportunity','Primary Opportunity','Lead Status','Assigned User'])
                    downloads[kind]['rows']=len(rows)
                else:
                    from openpyxl import load_workbook
                    workbook=load_workbook(io.BytesIO(response.content),read_only=True,data_only=True)
                    rows=list(workbook['ALL QUALIFIED'].values)
                    assert any('Website Opportunity' in row and 'Security Opportunity' in row for row in rows)
                    downloads[kind]['sheets']=workbook.sheetnames;workbook.close()
            call('POST',f'/api/leads/{lid}/report')
            response=call('GET',f'/api/leads/{lid}/report')
            assert 'Service opportunities' in response.text
            (output/'report.html').write_bytes(response.content)
            downloads['html']={'status':response.status_code,'bytes':len(response.content),'sandbox':response.headers.get('content-security-policy','')}
            state['verification']={'workflow':True,'assignment':True,'notes':True,'downloads':downloads,'scores_present':all(item.get('website_opportunity_score') is not None for item in leads['items'])}
    state_path.write_text(json.dumps(state,indent=2))
    print(json.dumps({'campaign':state['campaign_id'],'status':campaign['status'],'counters':campaign['counters'],'runs':campaign['runs']},indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('action',choices=['start','status','cancel','resume','verify']);parser.add_argument('--record',default=str(state_path));args=parser.parse_args();state_path=Path(args.record);main(args.action)
